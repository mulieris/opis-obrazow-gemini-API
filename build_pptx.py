#!/usr/bin/env python3
"""Generuje prezentacje Tab3 w formacie .pptx (czarne tlo + lime)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG      = RGBColor(0x08, 0x09, 0x0A)
PANEL   = RGBColor(0x16, 0x18, 0x1A)
PANEL2  = RGBColor(0x1F, 0x24, 0x18)
LIME    = RGBColor(0xC6, 0xFF, 0x3D)
LIME_DIM= RGBColor(0x9B, 0xC6, 0x2F)
TEXT    = RGBColor(0xF3, 0xF5, 0xF0)
MUTED   = RGBColor(0x9A, 0xA0, 0x9A)
WARN    = RGBColor(0xFF, 0xD2, 0x4A)
BAD     = RGBColor(0xFF, 0x6B, 0x6B)
CODEBG  = RGBColor(0x0B, 0x0D, 0x0E)
CODEFG  = RGBColor(0x9B, 0xE3, 0x6B)

SANS = "Segoe UI"
MONO = "Consolas"

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def box(s, l, t, w, h, fill=PANEL, line=None, line_w=0.75):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    r.adjustments[0] = 0.06
    r.fill.solid(); r.fill.fore_color.rgb = fill
    r.shadow.inherit = False
    if line is None:
        r.line.color.rgb = RGBColor(0x2A, 0x2E, 0x2A); r.line.width = Pt(0.75)
    else:
        r.line.color.rgb = line; r.line.width = Pt(line_w)
    return r


def txt(s, l, t, w, h, runs, size=18, color=TEXT, bold=False, font=SANS,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=4, line_spacing=1.1):
    tb = s.shapes.add_textbox(l, t, w, h); tf = tb.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    if isinstance(runs, str):
        runs = [[(runs, color, bold)]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space); p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, color, bold)]
        for (rtext, rcolor, rbold) in para:
            r = p.add_run(); r.text = rtext
            r.font.size = Pt(size); r.font.color.rgb = rcolor
            r.font.bold = rbold; r.font.name = font
    return tb


def frame_label(s, name):
    sq = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(0.45), Inches(0.13), Inches(0.13))
    sq.fill.solid(); sq.fill.fore_color.rgb = LIME; sq.line.fill.background(); sq.shadow.inherit = False
    txt(s, Inches(1.05), Inches(0.36), Inches(7), Inches(0.4),
        name, size=12, color=LIME_DIM, font=MONO)


def title(s, main, tag=None):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.85), Inches(1.0), Inches(0.08), Inches(0.62))
    bar.fill.solid(); bar.fill.fore_color.rgb = LIME; bar.line.fill.background(); bar.shadow.inherit = False
    txt(s, Inches(1.1), Inches(0.92), Inches(9.5), Inches(0.8), main, size=30, bold=True, color=TEXT)
    if tag:
        tb = box(s, Inches(10.0), Inches(1.02), Inches(2.5), Inches(0.5), fill=PANEL, line=RGBColor(0x2A,0x2E,0x2A))
        txt(s, Inches(10.0), Inches(1.04), Inches(2.5), Inches(0.46), tag, size=12, color=MUTED,
            font=MONO, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def signature(s):
    txt(s, Inches(0.85), Inches(7.02), Inches(11.6), Inches(0.4),
        [[("Tab3 · opis zdjęcia przez Gemini   ·   ", MUTED, False),
          ("Opracowała: Ewelina Kasprowicz", LIME_DIM, True)]],
        size=10, font=MONO)


def chip(s, l, t, w, label, fill=PANEL, color=MUTED):
    b = box(s, l, t, w, Inches(0.42), fill=fill, line=RGBColor(0x2A,0x2E,0x2A))
    txt(s, l, t, w, Inches(0.42), label, size=11, color=color, font=MONO,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def codebox(s, l, t, w, h, lines):
    accent = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, Inches(0.06), h)
    accent.fill.solid(); accent.fill.fore_color.rgb = LIME; accent.line.fill.background(); accent.shadow.inherit = False
    b = box(s, l, t, w, h, fill=CODEBG, line=RGBColor(0x22,0x26,0x22))
    tb = s.shapes.add_textbox(l + Inches(0.2), t + Inches(0.12), w - Inches(0.35), h - Inches(0.24))
    tf = tb.text_frame; tf.word_wrap = True
    for i, (line, col) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.15
        r = p.add_run(); r.text = line
        r.font.size = Pt(12.5); r.font.name = MONO; r.font.color.rgb = col
    return b


# ---------------------------------------------------------------- SLIDE 1
s = slide()
frame_label(s, "Tab3.frame · 1440 × 900")
chip(s, Inches(0.85), Inches(1.0), Inches(3.0), "Projekt zaliczeniowy", color=LIME_DIM)
chip(s, Inches(3.95), Inches(1.0), Inches(3.7), "Laboratorium programowania", color=LIME_DIM)
chip(s, Inches(7.75), Inches(1.0), Inches(2.3), "r.a. 2025/2026", color=LIME_DIM)
txt(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(1.8),
    [[("Tab3 — ", TEXT, True), ("„Opisz zdjęcie”", LIME, True)],
     [("z wykorzystaniem Gemini", TEXT, True)]],
    size=44, bold=True, line_spacing=1.02)
txt(s, Inches(0.85), Inches(3.7), Inches(10.5), Inches(1.0),
    "Jak działa zakładka, z jakich plików się składa, jaka jest idea — oraz co trzeba było naprawić, żeby całość rzeczywiście ruszyła.",
    size=17, color=MUTED, line_spacing=1.3)
for i, label in enumerate(["Vaadin · frontend", "Spring Boot · backend", "Google Gemini API", "REST + WebClient"]):
    chip(s, Inches(0.85 + i*3.0), Inches(4.95), Inches(2.85), label, fill=PANEL2, color=LIME)
txt(s, Inches(0.85), Inches(5.75), Inches(11), Inches(0.6),
    [[("Opracowała: ", LIME, True), ("Ewelina Kasprowicz", TEXT, True)]], size=20)
signature(s)

# ---------------------------------------------------------------- SLIDE 2
s = slide(); frame_label(s, "idea"); title(s, "Idea", "po co to jest")
txt(s, Inches(0.85), Inches(1.95), Inches(11.6), Inches(1.1),
    [[("Zakładka ", TEXT, False), ("Tab3", TEXT, True),
      (" pozwala przesłać ", TEXT, False), ("zdjęcie", TEXT, True),
      (" oraz wpisać ", TEXT, False), ("instrukcję tekstową", TEXT, True),
      (" (np. „Jaki to owoc i jakiego koloru?”). Aplikacja wysyła obraz + instrukcję do modelu ", TEXT, False),
      ("Gemini", LIME, True),
      (", a ten zwraca opis treści obrazu zgodny z instrukcją.", TEXT, False)]],
    size=18, line_spacing=1.35)
cards = [("Wejście", "Obraz (PNG/JPEG) + instrukcja tekstowa od użytkownika."),
         ("Przetwarzanie", "Gemini analizuje obraz w kontekście instrukcji i tworzy opis."),
         ("Wyjście", "JSON { found, answer } → opis pokazany w powiadomieniu.")]
for i, (h, p) in enumerate(cards):
    l = Inches(0.85 + i*4.05)
    box(s, l, Inches(3.6), Inches(3.8), Inches(2.0))
    txt(s, l + Inches(0.25), Inches(3.8), Inches(3.3), Inches(0.5), h, size=16, bold=True, color=LIME, font=MONO)
    txt(s, l + Inches(0.25), Inches(4.4), Inches(3.35), Inches(1.4), p, size=14, color=MUTED, line_spacing=1.3)
signature(s)

# ---------------------------------------------------------------- SLIDE 3
s = slide(); frame_label(s, "data-flow"); title(s, "Przepływ danych", "end-to-end")
steps = [("1 · Tab3.java", "Upload zdjęcia + pole tekstowe + przycisk „Opisz”."),
         ("2 · WebService", "POST /api/v1/gemini/photo przez WebClient."),
         ("3 · GeminiController", "Odbiera PhotoChangeRequest."),
         ("4 · GeminiService", "Buduje prompt, woła Gemini API."),
         ("5 · Gemini (Google)", "Analiza obrazu + instrukcji, odpowiedź JSON."),
         ("6 · GeminiParser", "Wyciąga tekst i mapuje na PhotoChangeResponse."),
         ("7 · → Tab3", "Mapowane na PhotoToChange → powiadomienie z opisem.")]
for i, (h, d) in enumerate(steps):
    col = i % 4; row = i // 4
    l = Inches(0.85 + col*3.05); t = Inches(2.1 + row*1.7)
    box(s, l, t, Inches(2.85), Inches(1.45))
    txt(s, l + Inches(0.18), t + Inches(0.12), Inches(2.5), Inches(0.4), h, size=12.5, color=LIME, font=MONO, bold=True)
    txt(s, l + Inches(0.18), t + Inches(0.58), Inches(2.55), Inches(0.8), d, size=12, color=MUTED, line_spacing=1.2)
signature(s)

# ---------------------------------------------------------------- SLIDE 4
s = slide(); frame_label(s, "components"); title(s, "Co dodaliśmy", "nowe pliki")
txt(s, Inches(0.85), Inches(1.85), Inches(5), Inches(0.4), "FRONTEND", size=12, color=MUTED, font=MONO)
txt(s, Inches(6.95), Inches(1.85), Inches(5), Inches(0.4), "BACKEND", size=12, color=MUTED, font=MONO)
fe = [("Tab3.java", "Widok: Upload, podgląd obrazu, pole instrukcji, przycisk."),
      ("tab3.css", "Stylizacja zakładki (układ, kolory, zaokrąglenia)."),
      ("WebService.isValidChange()", "Wysyła obraz+tekst na /photo."),
      ("PhotoToChange.java", "Model odpowiedzi frontu: found, answer.")]
be = [("GeminiController · /photo", "Endpoint REST: PhotoChangeRequest."),
      ("GeminiService.photoChange()", "Prompt do Gemini + obraz w base64."),
      ("PhotoChangeRequest / Response", "DTO żądania i odpowiedzi backendu."),
      ("GeminiParser.parsePhotoResponse()", "Parsowanie odpowiedzi Gemini.")]
for i, (h, p) in enumerate(fe):
    t = Inches(2.3 + i*1.12)
    box(s, Inches(0.85), t, Inches(5.7), Inches(0.98))
    txt(s, Inches(1.05), t + Inches(0.08), Inches(5.4), Inches(0.4),
        [[(h, LIME, True), ("  NEW", LIME_DIM, True)]], size=13, font=MONO)
    txt(s, Inches(1.05), t + Inches(0.5), Inches(5.4), Inches(0.4), p, size=11.5, color=MUTED)
for i, (h, p) in enumerate(be):
    t = Inches(2.3 + i*1.12)
    box(s, Inches(6.95), t, Inches(5.5), Inches(0.98))
    txt(s, Inches(7.15), t + Inches(0.08), Inches(5.2), Inches(0.4),
        [[(h, LIME, True), ("  NEW", LIME_DIM, True)]], size=13, font=MONO)
    txt(s, Inches(7.15), t + Inches(0.5), Inches(5.2), Inches(0.4), p, size=11.5, color=MUTED)
signature(s)

# ---------------------------------------------------------------- SLIDE 5
s = slide(); frame_label(s, "frontend"); title(s, "Frontend — Tab3.java", "widok")
txt(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(0.7),
    "Layout poziomy: po lewej karta z formularzem, po prawej podgląd zdjęcia. Przycisk aktywny dopiero gdy jest plik i tekst.",
    size=15, color=MUTED, line_spacing=1.3)
codebox(s, Inches(0.85), Inches(2.7), Inches(11.6), Inches(3.0), [
    ("// kliknięcie „Opisz” — wyślij obraz + instrukcję, pokaż wynik", RGBColor(0x5d,0x63,0x5c)),
    ("buttonPhotoChange.addClickListener(event -> {", CODEFG),
    ("    String text  = textFieldChange.getValue();", CODEFG),
    ("    byte[] image = memoryBufferChange.getInputStream().readAllBytes();", CODEFG),
    ("    WebService ws = ApplicationContext.getBean(WebService.class);", CODEFG),
    ("    PhotoToChange result = ws.isValidChange(text, image);", CODEFG),
    ("    String message = result.isFound() ? result.getAnswer()", CODEFG),
    ("                     : \"To nie jest coś co mogę opisać.\";", CODEFG),
    ("    new Notification(message, 15000).open();", CODEFG),
    ("});", CODEFG),
])
signature(s)

# ---------------------------------------------------------------- SLIDE 6
s = slide(); frame_label(s, "backend"); title(s, "Backend — prompt do Gemini", "photoChange()")
txt(s, Inches(0.85), Inches(1.85), Inches(11.6), Inches(0.7),
    "Obraz kodowany do base64 i wysyłany razem z instrukcją jako inline_data. Prompt wymusza odpowiedź w formacie JSON.",
    size=15, color=MUTED, line_spacing=1.3)
codebox(s, Inches(0.85), Inches(2.7), Inches(11.6), Inches(2.9), [
    ("String base64Image = Base64.getEncoder().encodeToString(image);", CODEFG),
    ("jsonPayload = \"\"\"", CODEFG),
    ("  { \"contents\":[{ \"parts\":[", CODEFG),
    ("    { \"text\":\"Opisz zawartość obrazu zgodnie z instrukcją: %s.\\n", CODEFG),
    ("       Zwróć JSON: { \\\"found\\\":true/false, \\\"answer\\\":\\\"opis\\\" }\" },", CODEFG),
    ("    { \"inline_data\":{ \"mime_type\":\"image/jpeg\", \"data\":\"%s\" } }", CODEFG),
    ("  ]}] }\"\"\".formatted(text, base64Image);", CODEFG),
    ("// odpowiedź → GeminiParser.parsePhotoResponse() → PhotoChangeResponse", RGBColor(0x5d,0x63,0x5c)),
])
txt(s, Inches(0.85), Inches(5.8), Inches(11.6), Inches(0.5),
    "Komunikacja przez java.net.http.HttpClient; model gemini-2.5-flash (multimodalny).",
    size=14, color=MUTED)
signature(s)

# ---------------------------------------------------------------- SLIDE 7
s = slide(); frame_label(s, "bug-report"); title(s, "Stan zastany", "dlaczego nie działało")
txt(s, Inches(0.85), Inches(1.9), Inches(11.6), Inches(0.6),
    "Po złożeniu kodu zakładka „nic nie robiła”. Diagnoza wykazała błędy blokujące kompilację i działanie backendu:",
    size=15, color=MUTED, line_spacing=1.3)
items = [("Backend się nie kompilował — pomieszane klasy i metody.", BAD),
         ("Brak endpointów /photo i /plate — żądania frontu trafiały w 404.", BAD),
         ("Niepoprawny JSON w promptcie (surowe znaki nowej linii).", BAD),
         ("Złe nazwy akcesorów — dane nie wiązały się z JSON-em.", BAD),
         ("Brak getterów w ImageResponse → błąd 406.", BAD),
         ("Puste catch {} — błędy ukryte, brak informacji.", WARN),
         ("Backend nie był uruchomiony — front pukał na pusty port 8081.", WARN)]
runs = [[("▸  ", c, True), (t, TEXT, False)] for (t, c) in items]
txt(s, Inches(0.85), Inches(2.7), Inches(11.6), Inches(4.0), runs, size=17, line_spacing=1.25, space=8)
signature(s)

# ---------------------------------------------------------------- SLIDES 8 & 9 (fixes)
def fixes_slide(name, tag, data):
    s = slide(); frame_label(s, name); title(s, "Co poprawiliśmy", tag)
    t0 = 1.95
    for (ic, head, body) in data:
        box(s, Inches(0.85), Inches(t0), Inches(11.6), Inches(1.12))
        txt(s, Inches(1.05), Inches(t0 + 0.1), Inches(11.2), Inches(0.45),
            [[(ic + "  ", TEXT, False), (head, TEXT, True), ("   FIX", WARN, True)]], size=15)
        txt(s, Inches(1.05), Inches(t0 + 0.55), Inches(11.2), Inches(0.55), body, size=12.5, color=MUTED, line_spacing=1.2)
        t0 += 1.28
    signature(s)

fixes_slide("fixes · 1/2", "żeby zadziałało · 1", [
    ("🧩", "PhotoChangeResponse.java", "Plik zawierał omyłkowo klasę GeminiParser (copy-paste). Przywróciliśmy właściwą klasę z polami found / answer — bez niej kod się nie kompilował."),
    ("🧹", "GeminiParser.java", "Zdublowane metody, literówka paresImageResponse, getText() bez return, brak parsePhotoResponse. Przepisaliśmy na czystą wersję z 4 metodami."),
    ("🔌", "GeminiController — dodane /photo i /plate", "Frontend wołał te ścieżki, ale ich nie było. Dodaliśmy @PostMapping wywołujące photoChange() i isCorrect()."),
])
fixes_slide("fixes · 2/2", "żeby zadziałało · 2", [
    ("📝", "GeminiService.photoChange() — naprawiony JSON", "Tekst promptu rozbity na 3 linie → niepoprawny JSON. Zwinęliśmy do jednej linii z \\n i poprawnymi cudzysłowami."),
    ("🔗", "PhotoChangeRequest — nazwa settera", "setSentenceToChange nie pasował do klucza JSON sentenseToChange → instrukcja docierała jako null. Zmiana na setSentenseToChange."),
    ("📦", "ImageResponse (backend) — dodane gettery", "Brak getterów → Jackson nie serializował odpowiedzi → 406. Dodaliśmy isFound() i getDetected_plate() (naprawia Tab2)."),
    ("🔔", "Tab3.java — widoczne błędy", "Puste catch {} zamieniliśmy na Notification.show(...) — błąd jest teraz widoczny."),
])

# ---------------------------------------------------------------- SLIDE 10
s = slide(); frame_label(s, "root-cause"); title(s, "Diagnoza: „kliknięcie nic nie robi”", "naprawdę")
txt(s, Inches(0.85), Inches(1.9), Inches(11.6), Inches(0.6),
    "Po naprawie kompilacji zakładka nadal nie reagowała. Sprawdzenie portów ujawniło przyczynę:",
    size=15, color=MUTED, line_spacing=1.3)
codebox(s, Inches(0.85), Inches(2.65), Inches(11.6), Inches(1.4), [
    ("frontend :8080  ->  200    // działał", CODEFG),
    ("backend  :8081  ->  000    // connection refused — NIE był uruchomiony!", WARN),
])
txt(s, Inches(0.85), Inches(4.3), Inches(11.6), Inches(1.0),
    "Frontend tylko przekazuje żądanie do backendu (port 8081), a ten dopiero rozmawia z Gemini. Bez uruchomionego backendu przycisk „Opisz” zawsze kończył się błędem.",
    size=16, line_spacing=1.35)
txt(s, Inches(0.85), Inches(5.6), Inches(11.6), Inches(0.7),
    [[("→ Po uruchomieniu backendu wszystkie 3 endpointy zwracają 200, a zakładki działają end-to-end.", LIME, True)]],
    size=16)
signature(s)

# ---------------------------------------------------------------- SLIDE 11
s = slide(); frame_label(s, "verify"); title(s, "Weryfikacja", "że działa")
cards = [("POST /photo", "Tab3 — opis zdjęcia"),
         ("POST /plate", "Tab2 — tablice rejestracyjne"),
         ("GET /{lang}/{sentence}", "GrammarCheckTab — gramatyka")]
for i, (h, p) in enumerate(cards):
    l = Inches(0.85 + i*4.05)
    box(s, l, Inches(2.0), Inches(3.8), Inches(2.0))
    txt(s, l + Inches(0.22), Inches(2.15), Inches(3.4), Inches(0.4), h, size=13.5, color=LIME, font=MONO, bold=True)
    txt(s, l + Inches(0.22), Inches(2.65), Inches(3.4), Inches(0.5), "200 ✓", size=26, color=LIME, bold=True)
    txt(s, l + Inches(0.22), Inches(3.35), Inches(3.4), Inches(0.5), p, size=13, color=MUTED)
txt(s, Inches(0.85), Inches(4.35), Inches(11), Inches(0.4), "Obie części budują się przez Maven bez błędów:", size=15, color=MUTED)
codebox(s, Inches(0.85), Inches(4.95), Inches(11.6), Inches(1.4), [
    ("sh mvnw -pl backend  -am compile    # BUILD SUCCESS", CODEFG),
    ("sh mvnw -pl frontend     compile    # BUILD SUCCESS", CODEFG),
    ("sh mvnw -pl backend spring-boot:run # start na :8081", CODEFG),
])
signature(s)

# ---------------------------------------------------------------- SLIDE 12
s = slide(); frame_label(s, "kryteria"); title(s, "Spełnione kryteria", "wymagania zaliczenia")
rows = [
    ("3", [("Część frontendowa aplikacji", "aplikacja Vaadin z zakładkami (gramatyka, tablice, opis zdjęcia)"),
           ("Poprawna struktura i czytelny interfejs", "moduły frontend/backend, pakiety, style CSS")]),
    ("4", [("Backend + co najmniej jeden endpoint API", "GeminiController: GET oraz POST /photo, /plate"),
           ("Komunikacja frontend ↔ backend", "WebService (WebClient) wysyła obraz+tekst, odbiera JSON")]),
    ("5", [("Rozwiązanie wykorzystujące Gemini", "GeminiService woła Gemini API w 3 funkcjach"),
           ("Projekt czytelny i możliwy do uruchomienia", "kompiluje się i działa end-to-end"),
           ("Dodatkowe funkcjonalności ponad minimum", "upload plików, podgląd obrazu, wybór języka")]),
]
t0 = 1.95
for grade, lst in rows:
    h = Inches(0.55 + 0.45*len(lst) + 0.25)
    box(s, Inches(0.85), Inches(t0), Inches(11.6), h)
    gb = box(s, Inches(1.05), Inches(t0 + 0.18), Inches(0.95), Inches(0.85), fill=PANEL2, line=LIME_DIM)
    txt(s, Inches(1.05), Inches(t0 + 0.16), Inches(0.95), Inches(0.55), grade, size=26, color=LIME, bold=True, align=PP_ALIGN.CENTER)
    txt(s, Inches(1.05), Inches(t0 + 0.66), Inches(0.95), Inches(0.3), "OCENA", size=8.5, color=LIME_DIM, font=MONO, align=PP_ALIGN.CENTER)
    runs = [[("✓  ", LIME, True), (req, TEXT, True), ("  →  " + how, MUTED, False)] for (req, how) in lst]
    txt(s, Inches(2.25), Inches(t0 + 0.18), Inches(10.0), h - Inches(0.3), runs, size=13.5, line_spacing=1.15, space=5)
    t0 += (0.55 + 0.45*len(lst) + 0.25) + 0.2
signature(s)

# ---------------------------------------------------------------- SLIDE 13
s = slide(); frame_label(s, "summary"); title(s, "Podsumowanie")
items = [
    [("Idea: ", LIME, True), ("opis zdjęcia wg instrukcji użytkownika z użyciem modelu Gemini.", TEXT, False)],
    [("Dodano: ", LIME, True), ("pełną ścieżkę Tab3 (UI + CSS + klient HTTP + endpoint + logika + parser + DTO).", TEXT, False)],
    [("Naprawiono: ", LIME, True), ("kompilację backendu, brakujące endpointy, JSON, wiązanie danych, serializację (406) i widoczność błędów.", TEXT, False)],
    [("Bonus: ", LIME, True), ("przy okazji naprawiona także zakładka Tab2 (tablice).", TEXT, False)],
    [("Status: ", LIME, True), ("wszystkie 3 funkcje działają end-to-end — kryteria zaliczenia spełnione.", TEXT, False)],
]
runs = [[("▸  ", LIME, True)] + it for it in items]
txt(s, Inches(0.85), Inches(2.0), Inches(11.6), Inches(3.4), runs, size=17, line_spacing=1.3, space=10)
txt(s, Inches(0.85), Inches(5.5), Inches(11.6), Inches(0.5),
    "Do rozważenia na przyszłość: klucz API w zmiennej środowiskowej zamiast w kodzie.", size=14, color=MUTED)
txt(s, Inches(0.85), Inches(6.2), Inches(11), Inches(0.5),
    [[("Opracowała: ", LIME, True), ("Ewelina Kasprowicz", TEXT, True)]], size=16)
signature(s)

def disable_spellcheck(presentation):
    """Wylacza czerwone podkreslenia pisowni: noProof + lang pl-PL na WSZYSTKICH
    wlasciwosciach tekstu (rPr, defRPr, endParaRPr) w kazdym slajdzie."""
    from pptx.oxml.ns import qn
    tags = {qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")}
    for sl in presentation.slides:
        for el in sl._element.iter():
            if el.tag in tags:
                el.set("lang", "pl-PL")
                el.set("noProof", "1")


disable_spellcheck(prs)
prs.save("tab3-prezentacja.pptx")
print("OK: tab3-prezentacja.pptx —", len(prs.slides._sldIdLst), "slajdów (bez podkreśleń pisowni)")
